from __future__ import annotations

from io import BytesIO

from docx import Document

from agent_mentor.rag.chunking import chunk_sections
from agent_mentor.rag.documents import DocumentParser


def test_markdown_parser_preserves_heading_path() -> None:
    sections = DocumentParser(200).parse(
        "guide.md", b"# RAG\nGrounded answers.\n## Citations\nShow sources."
    )

    assert sections[0].heading_path == ["RAG"]
    assert sections[1].heading_path == ["RAG", "Citations"]


def test_markdown_parser_ignores_headings_inside_fenced_code() -> None:
    sections = DocumentParser(max_pdf_pages=5).parse(
        "notes.md",
        b"# RAG\nbefore\n```python\n# not a heading\nprint('ok')\n```\nafter",
    )

    assert len(sections) == 1
    assert sections[0].heading_path == ["RAG"]
    assert "# not a heading" in sections[0].text


def test_markdown_parser_strips_utf8_bom_before_heading_detection() -> None:
    sections = DocumentParser(max_pdf_pages=5).parse(
        "notes.md",
        "\ufeff# 知识入库\n正文".encode(),
    )

    assert sections[0].heading_path == ["知识入库"]


def test_docx_parser_and_chunker() -> None:
    document = Document()
    document.add_heading("Agent", level=1)
    document.add_paragraph("A personal learning assistant should retain evidence.")
    buffer = BytesIO()
    document.save(buffer)

    chunks = chunk_sections(DocumentParser(200).parse("guide.docx", buffer.getvalue()), 30, 5)

    assert chunks[0].heading_path == ["Agent"]
    assert chunks[0].content


def test_markdown_table_block_type_is_carried_to_chunks() -> None:
    sections = DocumentParser(200).parse(
        "guide.md",
        b"# Metrics\n| metric | meaning |\n| - | - |\n| recall | retrieved coverage |",
    )

    chunks = chunk_sections(sections, 200, 20)

    assert sections[0].block_type == "table"
    assert chunks[0].block_type == "table"


def test_code_block_preserves_line_breaks_in_chunking() -> None:
    sections = DocumentParser(200).parse(
        "guide.md",
        b"# Example\n```python\nprint('agent')\nprint('rag')\n```",
    )

    chunks = chunk_sections(sections, 200, 20)

    assert sections[0].block_type == "code"
    assert chunks[0].block_type == "code"
    assert "print('agent')\nprint('rag')" in chunks[0].content
